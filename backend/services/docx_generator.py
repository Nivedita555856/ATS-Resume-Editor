from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt

from backend.models.schemas import ResumeData

BASE_DIR = Path(__file__).resolve().parents[1]
OUTPUT_DIR = BASE_DIR / "generated"


class DocxGeneratorService:
    def __init__(self) -> None:
        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    @staticmethod
    def _set_default_font(document: Document) -> None:
        styles = document.styles
        normal_style = styles["Normal"]
        normal_style.font.name = "Calibri"
        normal_style.font.size = Pt(11)

        if "Resume Heading" not in [s.name for s in styles]:
            heading_style = styles.add_style("Resume Heading", WD_STYLE_TYPE.PARAGRAPH)
            heading_style.font.name = "Calibri"
            heading_style.font.size = Pt(12)
            heading_style.font.bold = True

    def generate(self, resume: ResumeData, output_basename: str) -> Path:
        document = Document()
        self._set_default_font(document)

        name = document.add_paragraph(resume.personal_info.name)
        name.runs[0].bold = True
        name.runs[0].font.size = Pt(18)

        contact_parts = [resume.personal_info.email, resume.personal_info.phone]
        if resume.personal_info.location:
            contact_parts.append(resume.personal_info.location)
        if resume.personal_info.linkedin:
            contact_parts.append(resume.personal_info.linkedin)
        if resume.personal_info.github:
            contact_parts.append(resume.personal_info.github)
        document.add_paragraph(" | ".join(contact_parts))

        if resume.summary:
            document.add_paragraph("Professional Summary", style="Resume Heading")
            document.add_paragraph(resume.summary)

        document.add_paragraph("Skills", style="Resume Heading")
        document.add_paragraph(", ".join(resume.skills) if resume.skills else "N/A")

        if resume.certifications:
            document.add_paragraph("Certifications", style="Resume Heading")
            for cert in resume.certifications:
                document.add_paragraph(cert, style="List Bullet")

        document.add_paragraph("Experience", style="Resume Heading")
        for exp in resume.experience:
            p = document.add_paragraph(f"{exp.title} | {exp.company} ({exp.start_date} - {exp.end_date})")
            p.runs[0].bold = True
            for bullet in exp.bullets:
                document.add_paragraph(bullet, style="List Bullet")

        document.add_paragraph("Education", style="Resume Heading")
        for edu in resume.education:
            edu_line = f"{edu.degree}, {edu.institution} ({edu.year})"
            if edu.gpa:
                edu_line += f" | GPA: {edu.gpa}"
            document.add_paragraph(edu_line)
            if edu.relevant_courses:
                document.add_paragraph(f"Relevant Courses: {', '.join(edu.relevant_courses)}")

        if resume.projects:
            document.add_paragraph("Projects", style="Resume Heading")
            for project in resume.projects:
                header = project.name
                if project.link:
                    header += f" ({project.link})"
                project_head = document.add_paragraph(header)
                project_head.runs[0].bold = True
                document.add_paragraph(project.description)
                if project.tech_stack:
                    document.add_paragraph(f"Tech Stack: {', '.join(project.tech_stack)}")

        output_path = OUTPUT_DIR / f"{output_basename}.docx"
        document.save(str(output_path))
        return output_path


docx_generator_service = DocxGeneratorService()
